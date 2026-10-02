"""
Génère les valeurs de référence des tests à partir des classeurs Excel d'origine.

Chaque scénario de tests/fixtures est saisi dans une copie du classeur, LibreOffice recalcule
le classeur, puis les résultats sont écrits dans tests/fixtures/<scénario>.expected.json.
Le test Node (tests/engines.test.js) compare ensuite les moteurs JavaScript à ces valeurs.

Usage :
    python3 tests/excel/reference.py --pv chemin/plus_value.xlsx --succession chemin/succession.xlsx

Prérequis : openpyxl, LibreOffice (soffice) dans le PATH.
"""
import argparse
import datetime as dt
import glob
import json
import os
import shutil
import subprocess
import tempfile
import warnings

import openpyxl

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
FIXTURES = os.path.join(HERE, "..", "fixtures")


def date(s):
    return dt.datetime.strptime(s, "%Y-%m-%d") if s else None


def recalc(path, workdir):
    out = os.path.join(workdir, "out")
    os.makedirs(out, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, path],
        check=True, capture_output=True,
    )
    return openpyxl.load_workbook(os.path.join(out, os.path.basename(path)), data_only=True)


def clear(ws, rows, cols):
    for r in rows:
        for c in cols:
            ws.cell(r, c).value = None


def fill_pv(wb, d):
    ws = wb["Dossier"]
    ws["B3"] = date(d["dateCession"])
    ws["B4"] = d["prixVente"]
    ws["B5"] = d["fraisVendeurs"]
    ws["B6"] = "Oui" if d["bati"] else "Non"
    ws["B14"] = d.get("ageUsufruitier")
    ws["B15"] = d.get("cleUsufruitVente")
    ws["B16"] = d.get("cleUsufruitAcquisition")
    clear(ws, range(21, 27), range(1, 4))
    for i, v in enumerate(d["vendeurs"]):
        ws.cell(21 + i, 1).value = v["nom"]
        ws.cell(21 + i, 2).value = "Oui" if v.get("exonere") else "Non"
    ws["B30"] = d.get("valeurTerrain")
    ws["B31"] = d.get("valeurConstruction")
    o = wb["Origines"]
    clear(o, range(5, 35), range(1, 14))
    for r in range(5, 35):
        o.cell(r, 8).value = 1
    for i, x in enumerate(d["origines"]):
        r = 5 + i
        vals = [x["vendeur"], x.get("libelle"), x["nature"], date(x["dateAcquisition"]), x["droitRecu"],
                x.get("droitCede"), x["quotePart"], x["partPrix"], x["prix"], x["devise"],
                x.get("fraisReels"), x.get("travauxReels"), date(x.get("dateAchevement"))]
        for c, v in enumerate(vals, start=1):
            o.cell(r, c).value = v


def read_pv(wb, d):
    o = wb["Origines"]
    origines = []
    for i in range(len(d["origines"])):
        r = 5 + i
        origines.append({k: o.cell(r, c).value for k, c in {
            "prixCessionBrut": 17, "fraisCession": 18, "prixCessionNet": 19, "prixAcquisition": 20,
            "fraisAcquisition": 21, "annees": 22, "travaux": 23, "prixMajore": 24, "pvBrute": 25,
            "abatIR": 26, "abatPS": 27, "pvIR": 28, "pvPS": 29, "regle": 30, "controle": 31}.items()})
    rc = wb["Récapitulatif"]
    vendeurs = []
    for i in range(len(d["vendeurs"])):
        c = 2 + i
        vendeurs.append({k: rc.cell(r, c).value for k, r in {
            "prixCessionBrut": 7, "fraisCession": 8, "prixCessionNet": 9, "prixAcquisition": 10,
            "fraisAcquisition": 11, "travaux": 12, "prixMajore": 13, "pvBrute": 14, "pvImposableIR": 15,
            "pvImposablePS": 16, "impotRevenu": 17, "surtaxe": 18, "prelevementsSociaux": 19, "total": 20}.items()})
    return {"origines": origines, "vendeurs": vendeurs,
            "controleRepartition": rc["B22"].value, "controleLignes": rc["B23"].value}


def fill_succession(wb, d):
    ws = wb["Dossier"]
    ws["C6"] = d["nomDossier"]
    ws["C7"] = date(d["dateSimulation"])
    ws["C8"] = d["nbParents"]
    for i in range(2):
        p = d["parents"][i]
        ws.cell(12 + i, 3).value = p.get("nom") or None
        ws.cell(12 + i, 4).value = p.get("age")
    names = {}
    for base, key in ((17, "enfants"), (27, "petitsEnfants")):
        clear(ws, range(base, base + 6), range(2, 5))
        for i, e in enumerate(d[key]):
            ws.cell(base + i, 2).value = e["nom"]
            ws.cell(base + i, 3).value = "Oui" if e["majeur"] else "Non"
            ws.cell(base + i, 4).value = e.get("part")
            names[e["id"]] = e["nom"]
    ws["C36"] = "Oui" if d["donFamilial"] else "Non"
    ws["C37"] = "Oui" if d["deces15ans"] else "Non"
    p = wb["Patrimoine"]
    clear(p, range(6, 31), [2, 3, 4, 5, 7, 8, 9, 10, 11])
    for i, b in enumerate(d["biens"]):
        r = 6 + i
        vals = {2: b["designation"], 3: b["nature"], 4: b["valeur"], 5: b.get("quotePartP1"), 7: b["regime"],
                8: "Oui" if b["aDonner"] else "Non", 9: b.get("montantDonne"), 10: b["beneficiaire"], 11: b.get("montantPE")}
        for c, v in vals.items():
            p.cell(r, c).value = v
    a = wb["Antérieures"]
    clear(a, range(6, 26), range(1, 6))
    for i, x in enumerate(d["anterieures"]):
        r = 6 + i
        a.cell(r, 1).value = date(x["date"])
        a.cell(r, 2).value = x["montant"]
        a.cell(r, 3).value = "Parent %d" % x["donateur"]
        a.cell(r, 4).value = names[x["beneficiaire"]]


def read_block(ws, rows, cols):
    return [[ws.cell(r, c).value for c in cols] for r in rows]


def read_succession(wb, d):
    s, don, syn = wb["Succession"], wb["Donations"], wb["Synthèse"]
    ne, npe = len(d["enfants"]), len(d["petitsEnfants"])
    res = {
        # Colonnes C à I : transmis, abattement, exonérations, reste taxable, droits, taux, net.
        "A_cumul": read_block(s, range(11, 11 + ne), range(3, 10)),
        "A_p1": read_block(s, range(21, 21 + ne), range(3, 10)),
        "A_p2": read_block(s, range(32, 32 + ne), range(3, 10)),
    }
    # Colonnes C à N : donation (C-G), résiduel (H-L), total des droits (M), économie (N).
    for code, start in (("B", 31), ("C", 95)):
        for suffix, off in (("cumul", 0), ("p1", 20), ("p2", 40)):
            r0 = start + off
            res["%s_%s_enfants" % (code, suffix)] = read_block(don, range(r0, r0 + ne), range(3, 15))
            res["%s_%s_pe" % (code, suffix)] = read_block(don, range(r0 + 8, r0 + 8 + npe), range(3, 8))
            res["%s_%s_total" % (code, suffix)] = read_block(don, [r0 + 15], range(3, 15))[0]
    res["biensDonnes"] = [row for row in read_block(don, range(12, 22), range(2, 9)) if row[0] not in (None, "")]
    res["usufruit"] = read_block(don, range(7, 9), range(3, 6))
    res["comparatif"] = read_block(syn, range(13, 16), range(3, 9))
    res["synthese"] = {"patrimoine": syn["B7"].value, "sansDonation": syn["C7"].value, "meilleure": syn["E7"].value}
    res["controles"] = [wb["Dossier"].cell(r, 3).value for r in range(40, 46)]
    return res


def run(kind, src):
    fill, read = (fill_pv, read_pv) if kind == "plusvalue" else (fill_succession, read_succession)
    for fx in sorted(glob.glob(os.path.join(FIXTURES, kind + "-[0-9]*.json"))):
        if fx.endswith(".expected.json"):
            continue
        scenario = json.load(open(fx, encoding="utf-8"))
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, kind + ".xlsx")
            shutil.copy(src, path)
            wb = openpyxl.load_workbook(path)
            fill(wb, scenario["dossier"])
            wb.save(path)
            result = read(recalc(path, tmp), scenario["dossier"])
        out = fx.replace(".json", ".expected.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=1, default=str)
        print("écrit", os.path.relpath(out))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--pv")
    ap.add_argument("--succession")
    args = ap.parse_args()
    if args.pv:
        run("plusvalue", args.pv)
    if args.succession:
        run("succession", args.succession)
